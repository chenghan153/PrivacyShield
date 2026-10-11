#!/usr/bin/env python3
"""Build the 15-slide English Stage I proposal and convert it to PDF."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "slides"
ASSETS = ROOT / "demo" / "outputs"
RESULTS = json.loads((ROOT / "evaluation" / "results.json").read_text(encoding="utf-8"))
GITHUB = "https://github.com/chenghan153/PrivacyShield"
YOUTUBE = os.environ.get("YOUTUBE_URL", "TO ADD AFTER OWNER UPLOAD")

NAVY = RGBColor(16, 42, 67)
INK = RGBColor(23, 32, 51)
SLATE = RGBColor(71, 85, 105)
MUTED = RGBColor(100, 116, 139)
LIGHT = RGBColor(241, 245, 249)
LINE = RGBColor(203, 213, 225)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(20, 184, 166)
MINT = RGBColor(102, 227, 196)
BLUE = RGBColor(37, 99, 235)
AMBER = RGBColor(245, 158, 11)
RED = RGBColor(220, 38, 38)


def add_text(slide, x, y, w, h, text, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def box(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    return shape


def base(prs, title, kicker=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY if dark else RGBColor(248, 250, 252)
    if kicker:
        add_text(slide, 0.55, 0.22, 12.2, 0.3, kicker.upper(), 10, MINT if dark else TEAL, True)
    add_text(slide, 0.55, 0.48, 12.2, 0.62, title, 27, WHITE if dark else NAVY, True)
    if not dark:
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.12), Inches(0.7), Inches(0.055))
        accent.fill.solid()
        accent.fill.fore_color.rgb = TEAL
        accent.line.fill.background()
    add_text(slide, 12.15, 7.14, 0.65, 0.2, f"{len(prs.slides):02d}", 9, MINT if dark else MUTED, True, PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, body, accent=TEAL):
    box(slide, x, y, w, h)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, x + 0.25, y + 0.18, w - 0.45, 0.38, heading, 16, NAVY, True)
    add_text(slide, x + 0.25, y + 0.68, w - 0.45, h - 0.83, body, 12, SLATE)


def add_footer(slide, text="Stage I prototype · synthetic data · no UGen300/Hailo claim"):
    add_text(slide, 0.55, 7.13, 10.7, 0.2, text, 8, MUTED)


def hyperlink(shape, url):
    shape.click_action.hyperlink.address = url


def build() -> Path:
    OUT.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 — title
    slide = base(prs, "PrivacyShield", "ASUS UGen AI League · Workplace AI · Battlefield Lightning", dark=True)
    add_text(slide, 0.65, 1.65, 11.4, 1.25, "Audience-aware redaction\nbefore sensitive content is shared", 34, WHITE, True)
    add_text(slide, 0.67, 3.15, 9.7, 0.65, "Local-first · explainable policy · human-controlled · irreversible export", 18, MINT)
    box(slide, 0.67, 4.35, 4.2, 1.2, RGBColor(23, 58, 87), RGBColor(48, 84, 111))
    add_text(slide, 0.95, 4.58, 3.65, 0.7, "A working Stage I prototype\nusing only fictional demo data", 15, WHITE, True)
    add_text(slide, 0.67, 6.65, 11.5, 0.3, "2026 · Repository: github.com/chenghan153/PrivacyShield", 11, WHITE)

    # 2 — scenario
    slide = base(prs, "One screenshot. Three audiences. Different exposure.", "The sharing mistake")
    card(slide, 0.65, 1.45, 3.75, 4.75, "Internal team", "Project names and client context may be necessary.\n\nSecrets still never belong in a shared screenshot.", TEAL)
    card(slide, 4.78, 1.45, 3.75, 4.75, "External client", "Personal contact details, internal projects, and cost data can cross the wrong trust boundary.", AMBER)
    card(slide, 8.91, 1.45, 3.75, 4.75, "Public", "A post or presentation needs the strictest policy—and a missed field can be copied forever.", RED)
    add_text(slide, 0.65, 6.42, 11.7, 0.42, "Manual black boxes do not encode who the recipient is.", 20, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 3 — customer/problem
    slide = base(prs, "Built for teams that share work outside the company", "Target customer & problem")
    card(slide, 0.65, 1.5, 3.75, 2.15, "Who pays?", "SMB sales, finance, consulting, and outsourcing teams that routinely send screenshots and PDFs to clients or vendors.", TEAL)
    card(slide, 4.78, 1.5, 3.75, 2.15, "Why now?", "Cloud collaboration makes sharing instant, but manual redaction remains slow, inconsistent, and audience-blind.", BLUE)
    card(slide, 8.91, 1.5, 3.75, 2.15, "What is at risk?", "Credentials, personal contacts, internal codenames, client identity, and commercial terms.", RED)
    box(slide, 0.65, 4.05, 12.0, 2.1, NAVY, NAVY)
    add_text(slide, 1.0, 4.38, 11.3, 0.55, "Job to be done", 15, MINT, True)
    add_text(slide, 1.0, 4.95, 11.3, 0.75, "“Before I share this file, show me what this recipient should not see—and let me decide.”", 25, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 4 — alternatives
    slide = base(prs, "The gap is policy at the moment of sharing", "Existing workflow")
    headers = ["Approach", "Useful for", "Gap at sharing time"]
    rows = [
        ("Manual blur / black box", "One-off edits", "Easy to miss fields; policy lives in memory"),
        ("Regex/OCR utilities", "Finding known patterns", "Usually audience-blind; review UX varies"),
        ("Enterprise DLP", "Central controls", "Can be heavy for SMB ad-hoc screenshots"),
        ("PrivacyShield", "Recipient-specific copies", "Stage I prototype; needs broader validation"),
    ]
    xs = [0.7, 4.05, 7.7]
    ws = [3.15, 3.45, 4.9]
    for x, w, text_value in zip(xs, ws, headers):
        box(slide, x, 1.45, w, 0.62, NAVY, NAVY, False)
        add_text(slide, x + 0.12, 1.53, w - 0.24, 0.42, text_value, 13, WHITE, True)
    y = 2.08
    for i, row in enumerate(rows):
        fill = RGBColor(232, 250, 246) if i == 3 else WHITE
        for x, w, text_value in zip(xs, ws, row):
            box(slide, x, y, w, 0.92, fill, LINE, False)
            add_text(slide, x + 0.12, y + 0.08, w - 0.24, 0.72, text_value, 12, NAVY if i == 3 else SLATE, i == 3)
        y += 0.92
    add_text(slide, 0.7, 6.15, 11.8, 0.48, "Innovation claim: an integrated, audience-aware local workflow—not a claim that OCR, DLP, or contextual redaction is new.", 14, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, "Category comparison; no vendor capability claim")

    # 5 — actual before/after
    slide = base(prs, "From source document to recipient-safe copy", "Working product")
    slide.shapes.add_picture(str(ASSETS / "external_client_ui.png"), Inches(0.65), Inches(1.35), width=Inches(12.0))
    add_text(slide, 0.85, 6.6, 11.5, 0.35, "Scripted capture · image and masks generated by the real local engine · all data is fictional", 11, MUTED, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 6 — policy matrix
    slide = base(prs, "The same finding changes with the recipient", "Audience-aware policy")
    headers = ["Information", "Internal", "External client", "Public"]
    rows = [
        ("Email / phone", "ALLOW", "REDACT", "REDACT"),
        ("API key / secret", "REDACT", "REDACT", "REDACT"),
        ("Internal project", "ALLOW", "REDACT", "REDACT"),
        ("Client name", "ALLOW", "REVIEW", "REDACT"),
        ("Financial detail", "REVIEW", "REDACT", "REDACT"),
        ("Unknown", "REVIEW", "REVIEW", "REVIEW"),
    ]
    xs = [0.7, 4.35, 6.85, 9.75]
    ws = [3.45, 2.3, 2.7, 2.85]
    for x, w, label in zip(xs, ws, headers):
        box(slide, x, 1.35, w, 0.6, NAVY, NAVY, False)
        add_text(slide, x + 0.1, 1.44, w - 0.2, 0.4, label, 12, WHITE, True, PP_ALIGN.CENTER)
    y = 1.96
    colors = {"ALLOW": TEAL, "REDACT": RED, "REVIEW": AMBER}
    for row in rows:
        box(slide, xs[0], y, ws[0], 0.7, WHITE, LINE, False)
        add_text(slide, xs[0] + 0.18, y + 0.09, ws[0] - 0.3, 0.48, row[0], 12, INK, True)
        for x, w, value in zip(xs[1:], ws[1:], row[1:]):
            box(slide, x, y, w, 0.7, WHITE, LINE, False)
            add_text(slide, x, y + 0.09, w, 0.48, value, 11, colors[value], True, PP_ALIGN.CENTER)
        y += 0.7
    add_text(slide, 0.7, 6.35, 11.8, 0.45, "Demo defaults only. Administrators define production policy; uncertain content stays in human review.", 14, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 7 — journey
    slide = base(prs, "A six-step safety checkpoint before sharing", "User journey")
    labels = [
        ("1", "Upload", "PNG / JPG / PDF"),
        ("2", "Choose", "recipient audience"),
        ("3", "Detect", "OCR + rules"),
        ("4", "Explain", "type + reason"),
        ("5", "Review", "human checkboxes"),
        ("6", "Export", "fresh pixels"),
    ]
    for index, (number, heading, detail) in enumerate(labels):
        x = 0.63 + index * 2.08
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.55), Inches(1.65), Inches(0.8), Inches(0.8))
        circle.fill.solid(); circle.fill.fore_color.rgb = TEAL; circle.line.color.rgb = TEAL
        add_text(slide, x + 0.55, 1.67, 0.8, 0.72, number, 22, WHITE, True, PP_ALIGN.CENTER)
        if index < 5:
            connector = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + 1.52), Inches(1.88), Inches(0.38), Inches(0.32))
            connector.fill.solid(); connector.fill.fore_color.rgb = LINE; connector.line.color.rgb = LINE
        add_text(slide, x, 2.62, 1.9, 0.42, heading, 15, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, x, 3.08, 1.9, 0.55, detail, 11, SLATE, False, PP_ALIGN.CENTER)
    box(slide, 1.15, 4.4, 11.0, 1.45, NAVY, NAVY)
    add_text(slide, 1.45, 4.7, 10.4, 0.7, "Human review is a product feature—not an exception path.\nNo uncertain model suggestion alone can approve disclosure.", 20, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 8 — architecture
    slide = base(prs, "A deterministic safety core; acceleration comes later", "System architecture")
    implemented = ["Streamlit upload", "PDF text / Tesseract OCR", "Regex + dictionaries", "Audience policy", "Human review", "Raster PNG/PDF"]
    for index, value in enumerate(implemented):
        x = 0.7 + (index % 3) * 2.55
        y = 1.45 + (index // 3) * 1.35
        box(slide, x, y, 2.2, 0.85, WHITE, TEAL)
        add_text(slide, x + 0.1, y + 0.14, 2.0, 0.55, value, 12, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, 0.7, 4.28, 7.25, 0.35, "IMPLEMENTED & TESTED ON WORKSTATION CPU", 11, TEAL, True, PP_ALIGN.CENTER)
    box(slide, 8.55, 1.45, 4.05, 3.15, RGBColor(255, 251, 235), AMBER)
    add_text(slide, 8.88, 1.75, 3.4, 0.48, "Stage II validation plan", 18, NAVY, True)
    add_text(slide, 8.88, 2.38, 3.4, 1.55, "• Port OCR pre/post-processing\n• Benchmark Hailo-supported models\n• Measure latency and memory\n• Keep rules as fail-safe\n• Report only measured results", 14, SLATE)
    add_text(slide, 8.75, 4.86, 3.65, 0.55, "NOT YET TESTED ON\nUGen300 / Hailo", 15, RED, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 9 — local first
    slide = base(prs, "Local-first reduces a privacy paradox", "Why edge AI")
    card(slide, 0.65, 1.45, 3.75, 3.55, "Keep content local", "The prototype sends no document to an external AI API. Upload bytes stay in the local app process for the active session.", TEAL)
    card(slide, 4.78, 1.45, 3.75, 3.55, "Explain every action", "Deterministic categories map to readable reasons. Users can accept or reject every proposed mask.", BLUE)
    card(slide, 8.91, 1.45, 3.75, 3.55, "Destroy recoverable text", "Masks are burned into copied pixels. A fresh image-only PDF omits source text, annotations, and metadata.", AMBER)
    add_text(slide, 0.75, 5.45, 11.8, 0.78, "Local ≠ automatically secure. Host security, OCR misses, policy quality, and human review still matter.", 20, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 10 — three actual outputs
    slide = base(prs, "The policy difference is visible", "Actual engine outputs")
    images = [
        ("Internal Team", "internal_team_redacted.png", "2 selected"),
        ("External Client", "external_client_redacted.png", "6 selected"),
        ("Public", "public_redacted.png", "6 selected"),
    ]
    for index, (heading, file_name, detail) in enumerate(images):
        x = 0.62 + index * 4.24
        box(slide, x, 1.38, 3.95, 4.95, WHITE, LINE)
        add_text(slide, x + 0.2, 1.58, 3.55, 0.42, heading, 16, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, x + 0.2, 1.98, 3.55, 0.3, detail, 11, TEAL, True, PP_ALIGN.CENTER)
        slide.shapes.add_picture(str(ASSETS / file_name), Inches(x + 0.25), Inches(2.42), width=Inches(3.45), height=Inches(2.22))
        note = "Secret + budget review" if index == 0 else "Contact, project, secret, finance"
        add_text(slide, x + 0.2, 4.92, 3.55, 0.62, note, 11, SLATE, False, PP_ALIGN.CENTER)
    add_text(slide, 0.7, 6.55, 11.8, 0.28, "Actual outputs from demo/synthetic_dashboard.png · local OCR · fictional data", 10, MUTED, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 11 — evaluation
    det = RESULTS["detector"]
    bench = RESULTS["benchmark"]
    slide = base(prs, "Measured—and deliberately scoped", "Evaluation")
    metrics = [
        ("PYTEST", "10 passed", "unit + end-to-end"),
        ("SYNTHETIC FIXTURES", f"{det['tp']} / {det['fp']} / {det['fn']}", "TP / FP / FN"),
        ("SYSTEM MEDIAN", f"{bench['median_ms']:.2f} ms", "5 one-page PDFs"),
    ]
    for index, (k, value, detail) in enumerate(metrics):
        x = 0.75 + index * 4.15
        box(slide, x, 1.55, 3.7, 2.0, WHITE, TEAL)
        add_text(slide, x + 0.18, 1.8, 3.34, 0.35, k, 11, TEAL, True, PP_ALIGN.CENTER)
        add_text(slide, x + 0.18, 2.2, 3.34, 0.62, value, 25, NAVY, True, PP_ALIGN.CENTER)
        add_text(slide, x + 0.18, 2.86, 3.34, 0.35, detail, 11, MUTED, False, PP_ALIGN.CENTER)
    box(slide, 0.75, 4.05, 12.0, 1.65, RGBColor(255, 251, 235), AMBER)
    add_text(slide, 1.05, 4.32, 11.4, 0.95, "The 100% precision/recall fixture result covers 23 purpose-built text cases only. It bypasses OCR and is not a production accuracy claim. Human redaction time was not measured.", 17, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, 0.8, 6.05, 11.8, 0.48, f"Benchmark range {bench['min_ms']:.2f}–{bench['max_ms']:.2f} ms · workstation-specific · Public-policy raster export", 11, SLATE, False, PP_ALIGN.CENTER)
    add_footer(slide, "Source: evaluation/results.json generated 2026-10-11")

    # 12 — commercial hypothesis
    slide = base(prs, "Start with a team workflow, validate willingness to pay", "Commercial potential")
    card(slide, 0.65, 1.4, 3.6, 4.35, "Buyer", "Operations or IT lead at an SMB that shares client-facing screenshots and documents.\n\nEntry motion: self-serve team pilot.\n\nExpansion: centrally managed policies and audit controls.", TEAL)
    card(slide, 4.55, 1.4, 3.6, 4.35, "Pricing hypothesis", "$12 / user / month\nor $120 / small team / month\n\nTo validate through interviews; not a proven price or forecast.", BLUE)
    box(slide, 8.45, 1.4, 4.2, 4.35, WHITE, AMBER)
    add_text(slide, 8.72, 1.67, 3.65, 0.4, "Adjustable ROI scenario", 16, NAVY, True)
    add_text(slide, 8.72, 2.25, 3.65, 1.55, "500 docs/month\n× (3.0 − 0.25) min\n÷ 60 × $30/hour\n= $687.50 gross time value", 16, SLATE, True, PP_ALIGN.CENTER)
    add_text(slide, 8.72, 4.03, 3.65, 0.5, "− $120 team plan\n= $567.50 scenario value", 15, TEAL, True, PP_ALIGN.CENTER)
    add_text(slide, 8.72, 4.9, 3.65, 0.45, "All inputs are assumptions.\nNo manual timing study yet.", 11, RED, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 13 — product position
    slide = base(prs, "A focused checkpoint between creation and distribution", "Product position")
    card(slide, 0.65, 1.45, 3.72, 4.65, "Manual editors", "Strength\nFamiliar and flexible\n\nPrivacyShield wedge\nRecipient policy, finding reasons, repeatable export verification", AMBER)
    card(slide, 4.8, 1.45, 3.72, 4.65, "Cloud AI redaction", "Strength\nPotentially broader semantics\n\nPrivacyShield wedge\nNo document upload to an external AI service in the core workflow", BLUE)
    card(slide, 8.95, 1.45, 3.72, 4.65, "Enterprise DLP", "Strength\nCentral governance and coverage\n\nPrivacyShield wedge\nLightweight, preview-first flow for SMB ad-hoc sharing", TEAL)
    add_text(slide, 0.65, 6.37, 12.0, 0.4, "Next validation: 10–15 customer interviews + task-based redaction study on representative documents.", 15, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, "Positioning hypotheses; validate against named competitors before procurement claims")

    # 14 — risk and roadmap
    slide = base(prs, "Win trust by being explicit about failure modes", "Risks & roadmap")
    risks = [
        ("OCR miss", "Preview + confidence + broader test corpus"),
        ("Unknown secret format", "Admin dictionaries + rule packs + review"),
        ("Over-redaction", "Per-item control + recipient policy"),
        ("False sense of security", "Warnings + export verification + training"),
    ]
    for index, (risk, mitigation) in enumerate(risks):
        y = 1.36 + index * 1.1
        box(slide, 0.7, y, 5.7, 0.85, WHITE, LINE)
        add_text(slide, 0.92, y + 0.12, 1.65, 0.55, risk, 13, RED, True)
        add_text(slide, 2.55, y + 0.12, 3.55, 0.55, mitigation, 12, SLATE)
    box(slide, 6.82, 1.36, 5.8, 4.65, NAVY, NAVY)
    add_text(slide, 7.15, 1.67, 5.15, 0.4, "Roadmap", 18, MINT, True)
    add_text(slide, 7.15, 2.27, 5.15, 2.9, "NOW · Working CPU prototype\n\nNEXT · Real document study, policy admin, accessibility\n\nSTAGE II · UGen300/Hailo compatibility + measured acceleration\n\nLATER · Team controls and governed audit events", 15, WHITE)
    add_text(slide, 0.7, 6.35, 11.9, 0.42, "Go/no-go evidence for Stage II: measured device compatibility, latency, memory, and accuracy—not architecture diagrams.", 14, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 15 — links (course requirement)
    slide = base(prs, "Try the code. Watch the real workflow.", "Links & call to action", dark=True)
    add_text(slide, 0.75, 1.55, 11.8, 0.72, "PrivacyShield makes the recipient part of the redaction policy.", 27, WHITE, True, PP_ALIGN.CENTER)
    repo_box = box(slide, 1.05, 2.78, 11.2, 1.05, RGBColor(23, 58, 87), RGBColor(48, 84, 111))
    repo_text = add_text(slide, 1.35, 2.93, 10.6, 0.72, f"GitHub Repository\n{GITHUB}", 16, WHITE, True, PP_ALIGN.CENTER)
    hyperlink(repo_box, GITHUB); hyperlink(repo_text, GITHUB)
    video_box = box(slide, 1.05, 4.1, 11.2, 1.05, RGBColor(23, 58, 87), RGBColor(48, 84, 111))
    video_color = MINT if YOUTUBE.startswith("http") else RGBColor(253, 186, 116)
    video_text = add_text(slide, 1.35, 4.25, 10.6, 0.72, f"Demo Video\n{YOUTUBE}", 16, video_color, True, PP_ALIGN.CENTER)
    if YOUTUBE.startswith("http"):
        hyperlink(video_box, YOUTUBE); hyperlink(video_text, YOUTUBE)
    add_text(slide, 0.85, 5.72, 11.6, 0.55, "Stage I: local workstation CPU · Stage II plan: validate on ASUS UGen300 / Hailo", 14, WHITE, False, PP_ALIGN.CENTER)
    add_text(slide, 0.85, 6.42, 11.6, 0.42, "Thank you", 22, MINT, True, PP_ALIGN.CENTER)

    pptx = OUT / "PrivacyShield_Stage_I.pptx"
    prs.save(pptx)
    return pptx


def convert_pdf(pptx: Path) -> Path:
    libreoffice = shutil.which("libreoffice")
    if not libreoffice:
        raise RuntimeError("LibreOffice is required for PDF conversion")
    profile = ROOT / ".tmp" / "libreoffice-profile"
    profile.mkdir(parents=True, exist_ok=True)
    command = [
        libreoffice,
        "--headless",
        f"-env:UserInstallation=file://{profile}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(OUT),
        str(pptx),
    ]
    result = subprocess.run(command, check=False, text=True, capture_output=True, timeout=120)
    pdf = OUT / "PrivacyShield_Stage_I.pdf"
    if result.returncode != 0 or not pdf.exists():
        raise RuntimeError(f"LibreOffice conversion failed: {result.stdout}\n{result.stderr}")
    return pdf


if __name__ == "__main__":
    built = build()
    converted = convert_pdf(built)
    print(f"Built {built}")
    print(f"Built {converted}")
    print(f"YouTube field: {YOUTUBE}")
